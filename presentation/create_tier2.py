"""Generate Why 3 Tiers + Tier 2 Deep Learning presentation (NUS ISS template)."""
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from nus_iss_template import (
    create_presentation, add_title_slide, add_content_slide, add_section_slide,
    add_textbox, add_bullet_list, add_rect, add_rounded_rect,
    add_multiline_textbox,
    NUS_BLUE, NUS_ORANGE, WHITE, LIGHT_GRAY, SUBTLE_GRAY,
    ACCENT_BLUE, ACCENT_GREEN, ACCENT_PURPLE, ACCENT_RED,
    set_slide_bg, add_nus_logo, add_nus_footer,
    LIGHT_BLUE_BG, LIGHT_GREEN_BG, LIGHT_PURPLE_BG, LIGHT_ORANGE_BG,
    LIGHT_RED_BG, CARD_BG, LIGHT_BG,
)

TIER1_COLOR = ACCENT_BLUE
TIER2_COLOR = ACCENT_PURPLE
TIER3_COLOR = ACCENT_GREEN
DARK_CARD = RGBColor(0x00, 0x2A, 0x55)

prs = create_presentation()


def add_arrow_connector(slide, start_left, start_top, end_left, end_top,
                        color=SUBTLE_GRAY, width=Pt(2)):
    """Add a straight connector line (arrow)."""
    connector = slide.shapes.add_connector(
        1, start_left, start_top, end_left, end_top
    )
    connector.line.color.rgb = color
    connector.line.width = width
    return connector


# ============================================================
# SLIDE 1: Title
# ============================================================
add_title_slide(
    prs,
    title="Why 3 Tiers? + Tier 2: Deep Learning",
    subtitle="From classical baselines to neural networks to combined power",
    author="Muneeswaran Muthaiah",
    affiliation="Video2Knowledge  |  Pattern Recognition Systems",
    date_text="National University of Singapore",
)


# ============================================================
# SLIDE 2: Why 3 Tiers?
# ============================================================
slide = add_content_slide(prs, "Why Do We Need 3 Tiers?")

# --- Tier boxes ---
tier_box_w = Inches(3.5)
tier_box_h = Inches(3.2)
tier_y = Inches(1.3)
gap = Inches(0.4)
start_x = Inches(1.3)

# Tier 1
add_rounded_rect(slide, start_x, tier_y, tier_box_w, Inches(0.6),
                 TIER1_COLOR, "Tier 1: Classical ML", font_size=18,
                 font_color=WHITE, bold=True)
add_rounded_rect(slide, start_x, tier_y + Inches(0.6), tier_box_w,
                 tier_box_h - Inches(0.6), DARK_CARD)
add_multiline_textbox(slide, start_x + Inches(0.2), tier_y + Inches(0.75),
                      tier_box_w - Inches(0.4), tier_box_h - Inches(1.0),
                      [
                          {"text": "Fast baselines", "size": 15, "color": TIER1_COLOR, "bold": True},
                          {"text": "Establish performance floor", "size": 14, "color": NUS_ORANGE},
                          {"text": "", "size": 8, "color": NUS_ORANGE},
                          {"text": "Simple, interpretable,", "size": 14, "color": NUS_ORANGE},
                          {"text": "quick to train", "size": 14, "color": NUS_ORANGE},
                      ],
                      spacing=Pt(4))

# Tier 2
t2_x = start_x + tier_box_w + gap
add_rounded_rect(slide, t2_x, tier_y, tier_box_w, Inches(0.6),
                 TIER2_COLOR, "Tier 2: Deep Learning", font_size=18,
                 font_color=WHITE, bold=True)
add_rounded_rect(slide, t2_x, tier_y + Inches(0.6), tier_box_w,
                 tier_box_h - Inches(0.6), DARK_CARD)
add_multiline_textbox(slide, t2_x + Inches(0.2), tier_y + Inches(0.75),
                      tier_box_w - Inches(0.4), tier_box_h - Inches(1.0),
                      [
                          {"text": "Learn complex patterns", "size": 15, "color": TIER2_COLOR, "bold": True},
                          {"text": "automatically", "size": 15, "color": TIER2_COLOR, "bold": True},
                          {"text": "", "size": 8, "color": NUS_ORANGE},
                          {"text": "Can capture non-linear", "size": 14, "color": NUS_ORANGE},
                          {"text": "relationships between features", "size": 14, "color": NUS_ORANGE},
                      ],
                      spacing=Pt(4))

# Tier 3
t3_x = t2_x + tier_box_w + gap
add_rounded_rect(slide, t3_x, tier_y, tier_box_w, Inches(0.6),
                 TIER3_COLOR, "Tier 3: Ensemble", font_size=18,
                 font_color=WHITE, bold=True)
add_rounded_rect(slide, t3_x, tier_y + Inches(0.6), tier_box_w,
                 tier_box_h - Inches(0.6), DARK_CARD)
add_multiline_textbox(slide, t3_x + Inches(0.2), tier_y + Inches(0.75),
                      tier_box_w - Inches(0.4), tier_box_h - Inches(1.0),
                      [
                          {"text": "Combine the best of both", "size": 15, "color": TIER3_COLOR, "bold": True},
                          {"text": "", "size": 8, "color": NUS_ORANGE},
                          {"text": "Different models make different", "size": 14, "color": NUS_ORANGE},
                          {"text": "mistakes — combining them", "size": 14, "color": NUS_ORANGE},
                          {"text": "reduces overall error", "size": 14, "color": NUS_ORANGE},
                      ],
                      spacing=Pt(4))

# Arrows between tiers
arrow_y = tier_y + Inches(1.6)
add_arrow_connector(slide, start_x + tier_box_w, arrow_y,
                    t2_x, arrow_y, color=LIGHT_GRAY, width=Pt(3))
add_arrow_connector(slide, t2_x + tier_box_w, arrow_y,
                    t3_x, arrow_y, color=LIGHT_GRAY, width=Pt(3))

# Key message box at the bottom
key_msg_y = Inches(4.9)
add_rounded_rect(slide, Inches(1.5), key_msg_y, Inches(10.5), Inches(1.0),
                 DARK_CARD, font_size=16, font_color=WHITE, bold=False)
add_textbox(slide, Inches(1.9), key_msg_y + Inches(0.15), Inches(9.7), Inches(0.7),
            "No single model is best for everything. 3 tiers lets us compare "
            "approaches and pick the best for each use case.",
            font_size=17, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)


# ============================================================
# SLIDE 3: How Tiers Interconnect
# ============================================================
slide = add_content_slide(prs, "How Tiers Interconnect")

# 150-dim vector box (top center)
vec_x = Inches(5.2)
vec_y = Inches(1.3)
vec_w = Inches(3.0)
vec_h = Inches(0.6)
add_rounded_rect(slide, vec_x, vec_y, vec_w, vec_h,
                 NUS_ORANGE, "150-dim Feature Vector",
                 font_size=15, font_color=WHITE, bold=True)

# Tier 1 block (left)
t1_x = Inches(1.3)
t1_y = Inches(2.6)
t1_w = Inches(3.5)
t1_h = Inches(1.5)
add_rounded_rect(slide, t1_x, t1_y, t1_w, Inches(0.5),
                 TIER1_COLOR, "Tier 1: Classical ML",
                 font_size=15, font_color=WHITE, bold=True)
add_rounded_rect(slide, t1_x, t1_y + Inches(0.5), t1_w, t1_h - Inches(0.5),
                 DARK_CARD)
add_multiline_textbox(slide, t1_x + Inches(0.15), t1_y + Inches(0.55),
                      t1_w - Inches(0.3), Inches(0.9),
                      [
                          "SVM, Random Forest,",
                          "Logistic Regression, KNN, MLP",
                      ],
                      font_size=13, color=NUS_ORANGE)

# Tier 2 block (right)
t2_x = Inches(8.8)
t2_y = Inches(2.6)
t2_w = Inches(3.5)
t2_h = Inches(1.5)
add_rounded_rect(slide, t2_x, t2_y, t2_w, Inches(0.5),
                 TIER2_COLOR, "Tier 2: Deep Learning",
                 font_size=15, font_color=WHITE, bold=True)
add_rounded_rect(slide, t2_x, t2_y + Inches(0.5), t2_w, t2_h - Inches(0.5),
                 DARK_CARD)
add_multiline_textbox(slide, t2_x + Inches(0.15), t2_y + Inches(0.55),
                      t2_w - Inches(0.3), Inches(0.9),
                      [
                          "MLP, CNN-1D,",
                          "LSTM, Transformer",
                      ],
                      font_size=13, color=NUS_ORANGE)

# Arrows from vector to Tier 1 and Tier 2
add_arrow_connector(slide, vec_x, vec_y + vec_h,
                    t1_x + t1_w / 2, t1_y,
                    color=NUS_ORANGE, width=Pt(2.5))
add_arrow_connector(slide, vec_x + vec_w, vec_y + vec_h,
                    t2_x + t2_w / 2, t2_y,
                    color=NUS_ORANGE, width=Pt(2.5))

# "Predictions" labels
add_textbox(slide, t1_x + Inches(1.0), t1_y + t1_h + Inches(0.0),
            Inches(1.5), Inches(0.35),
            "predictions", font_size=12, color=LIGHT_GRAY,
            alignment=PP_ALIGN.CENTER)
add_textbox(slide, t2_x + Inches(1.0), t2_y + t2_h + Inches(0.0),
            Inches(1.5), Inches(0.35),
            "predictions", font_size=12, color=LIGHT_GRAY,
            alignment=PP_ALIGN.CENTER)

# Tier 3 block (center bottom)
t3_x = Inches(4.4)
t3_y = Inches(4.7)
t3_w = Inches(4.5)
t3_h = Inches(1.2)
add_rounded_rect(slide, t3_x, t3_y, t3_w, Inches(0.5),
                 TIER3_COLOR, "Tier 3: Ensemble Methods",
                 font_size=15, font_color=WHITE, bold=True)
add_rounded_rect(slide, t3_x, t3_y + Inches(0.5), t3_w, t3_h - Inches(0.5),
                 DARK_CARD)
add_textbox(slide, t3_x + Inches(0.15), t3_y + Inches(0.55),
            t3_w - Inches(0.3), Inches(0.6),
            "Combines Tier 1 + Tier 2 outputs",
            font_size=13, color=NUS_ORANGE, alignment=PP_ALIGN.CENTER)

# Arrows from Tier 1 and Tier 2 to Tier 3
add_arrow_connector(slide, t1_x + t1_w / 2, t1_y + t1_h,
                    t3_x, t3_y + Inches(0.3),
                    color=TIER1_COLOR, width=Pt(2.5))
add_arrow_connector(slide, t2_x + t2_w / 2, t2_y + t2_h,
                    t3_x + t3_w, t3_y + Inches(0.3),
                    color=TIER2_COLOR, width=Pt(2.5))

# Final prediction arrow
add_textbox(slide, t3_x + Inches(1.0), t3_y + t3_h + Inches(0.05),
            Inches(2.5), Inches(0.35),
            "Final Prediction", font_size=14, color=TIER3_COLOR,
            bold=True, alignment=PP_ALIGN.CENTER)

# Tier 3 uses Tier 1 and Tier 2 as building blocks
add_textbox(slide, Inches(1.3), Inches(6.3), Inches(5.5), Inches(0.4),
            "Tier 3 uses Tier 1 and Tier 2 as building blocks",
            font_size=15, color=LIGHT_GRAY, bold=True)

# Specific connections (right side)
conn_x = Inches(8.5)
conn_y = Inches(4.7)
add_textbox(slide, conn_x, conn_y, Inches(4.5), Inches(0.35),
            "Specific Ensemble Methods:", font_size=14,
            color=WHITE, bold=True)
add_multiline_textbox(slide, conn_x, conn_y + Inches(0.35),
                      Inches(4.5), Inches(1.8),
                      [
                          {"text": "Voting: SVM + RF + MLP", "size": 13, "color": NUS_ORANGE},
                          {"text": "Stacking: SVM + RF + MLP → LogReg", "size": 13, "color": NUS_ORANGE},
                          {"text": "Late Fusion: SVM (text) + RF (visual)", "size": 13, "color": NUS_ORANGE},
                      ],
                      spacing=Pt(6))


# ============================================================
# SLIDE 4: Tier 2 Overview
# ============================================================
slide = add_content_slide(prs, "Tier 2: Deep Learning Overview")

add_textbox(slide, Inches(1.3), Inches(1.3), Inches(11), Inches(0.5),
            "4 Neural Network architectures, each designed for different pattern types",
            font_size=20, color=LIGHT_GRAY)

# Four model cards
card_w = Inches(2.7)
card_h = Inches(3.5)
card_y = Inches(2.1)
card_gap = Inches(0.35)
card_start_x = Inches(1.3)

models = [
    ("MLP", "Multi-Layer\nPerceptron", ACCENT_BLUE,
     "Fully connected layers\nlearn feature combinations",
     "Best for: General\npattern matching"),
    ("CNN-1D", "1D Convolutional\nNeural Network", ACCENT_GREEN,
     "Sliding filters detect\nlocal feature patterns",
     "Best for: Adjacent\nfeature correlations"),
    ("LSTM", "Long Short-Term\nMemory", NUS_ORANGE,
     "Sequential processing\nwith memory gates",
     "Best for: Cross-modality\nlong-range dependencies"),
    ("Transformer", "Self-Attention\nNetwork", ACCENT_PURPLE,
     "Attention mechanism\nconnects any features",
     "Best for: Global feature\nrelationships"),
]

for i, (abbr, full_name, color, desc, best_for) in enumerate(models):
    cx = card_start_x + i * (card_w + card_gap)

    # Header
    add_rounded_rect(slide, cx, card_y, card_w, Inches(0.55),
                     color, abbr, font_size=20, font_color=WHITE, bold=True)

    # Body background
    add_rounded_rect(slide, cx, card_y + Inches(0.55), card_w,
                     card_h - Inches(0.55), DARK_CARD)

    # Full name
    add_textbox(slide, cx + Inches(0.15), card_y + Inches(0.7),
                card_w - Inches(0.3), Inches(0.7),
                full_name, font_size=13, color=color, bold=True,
                alignment=PP_ALIGN.CENTER)

    # Description
    add_textbox(slide, cx + Inches(0.15), card_y + Inches(1.4),
                card_w - Inches(0.3), Inches(0.8),
                desc, font_size=12, color=NUS_ORANGE,
                alignment=PP_ALIGN.CENTER)

    # Best for
    add_textbox(slide, cx + Inches(0.15), card_y + Inches(2.3),
                card_w - Inches(0.3), Inches(0.7),
                best_for, font_size=12, color=LIGHT_GRAY,
                alignment=PP_ALIGN.CENTER)

# Common tech note
add_rounded_rect(slide, Inches(1.5), Inches(5.9), Inches(10.5), Inches(0.7),
                 ACCENT_PURPLE)
add_textbox(slide, Inches(1.9), Inches(6.0), Inches(9.7), Inches(0.5),
            "All use PyTorch  |  Trained with gradient descent  |  GPU acceleration (CUDA)",
            font_size=16, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)


# ============================================================
# SLIDE 5: MLP (Multi-Layer Perceptron)
# ============================================================
slide = add_content_slide(prs, "MLP (Multi-Layer Perceptron)",
                          subtitle="The simplest neural network — fully connected layers")

# Architecture diagram (left side)
arch_x = Inches(1.3)
arch_y = Inches(1.3)

add_textbox(slide, arch_x, arch_y, Inches(5), Inches(0.4),
            "Architecture", font_size=18, color=WHITE, bold=True)

# Layer boxes
layer_x = Inches(1.5)
layer_w = Inches(4.5)
layer_h = Inches(0.55)
layer_gap = Inches(0.15)

layers = [
    ("Input: 150 features", NUS_ORANGE),
    ("Layer 1: 150 → 256 neurons + ReLU", ACCENT_BLUE),
    ("Layer 2: 256 → 128 neurons + ReLU", ACCENT_BLUE),
    ("Layer 3: 128 → 9 neurons + Softmax", ACCENT_PURPLE),
    ("Output: probability per class", ACCENT_GREEN),
]
ly = arch_y + Inches(0.5)
for label, col in layers:
    add_rounded_rect(slide, layer_x, ly, layer_w, layer_h,
                     col, label, font_size=14, font_color=WHITE, bold=False)
    ly += layer_h + layer_gap

# Arrows between layers
for i in range(len(layers) - 1):
    ay = arch_y + Inches(0.5) + (i + 1) * (layer_h + layer_gap) - layer_gap / 2
    add_textbox(slide, layer_x + Inches(2.0), ay - Inches(0.15),
                Inches(0.5), Inches(0.3),
                "↓", font_size=16, color=LIGHT_GRAY,
                alignment=PP_ALIGN.CENTER)

# Right side: explanations
exp_x = Inches(6.8)
exp_y = Inches(1.3)

add_textbox(slide, exp_x, exp_y, Inches(6), Inches(0.4),
            "Key Concepts", font_size=18, color=WHITE, bold=True)

# ReLU explanation
add_rounded_rect(slide, exp_x, exp_y + Inches(0.5), Inches(5.8), Inches(1.0),
                 DARK_CARD)
add_textbox(slide, exp_x + Inches(0.15), exp_y + Inches(0.55),
            Inches(5.5), Inches(0.3),
            "ReLU (Rectified Linear Unit)", font_size=15, color=ACCENT_BLUE,
            bold=True)
add_textbox(slide, exp_x + Inches(0.15), exp_y + Inches(0.85),
            Inches(5.5), Inches(0.6),
            "If negative → 0. Adds non-linearity so the network can learn complex patterns.",
            font_size=13, color=NUS_ORANGE)

# Softmax explanation
add_rounded_rect(slide, exp_x, exp_y + Inches(1.7), Inches(5.8), Inches(1.0),
                 DARK_CARD)
add_textbox(slide, exp_x + Inches(0.15), exp_y + Inches(1.75),
            Inches(5.5), Inches(0.3),
            "Softmax", font_size=15, color=ACCENT_PURPLE, bold=True)
add_textbox(slide, exp_x + Inches(0.15), exp_y + Inches(2.05),
            Inches(5.5), Inches(0.6),
            "Converts raw scores to probabilities summing to 1.",
            font_size=13, color=NUS_ORANGE)

# Strengths
add_textbox(slide, exp_x, exp_y + Inches(3.0), Inches(6), Inches(0.35),
            "Strengths", font_size=16, color=ACCENT_GREEN, bold=True)
add_multiline_textbox(slide, exp_x + Inches(0.15), exp_y + Inches(3.35),
                      Inches(5.5), Inches(0.8),
                      [
                          "•  Simple architecture, easy to understand",
                          "•  Fast inference (4ms per prediction)",
                          "•  Good baseline for comparison",
                      ],
                      font_size=13, color=NUS_ORANGE, spacing=Pt(4))

# What it learns (bottom bar)
add_rounded_rect(slide, Inches(1.3), Inches(5.7), Inches(11.3), Inches(0.8),
                 ACCENT_BLUE)
add_textbox(slide, Inches(1.7), Inches(5.8), Inches(10.5), Inches(0.6),
            'What it learns: "Combinations of features — e.g., high OCR confidence '
            '+ high interaction rate = coding activity"',
            font_size=15, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)


# ============================================================
# SLIDE 6: CNN-1D
# ============================================================
slide = add_content_slide(prs, "CNN-1D (1D Convolutional Neural Network)",
                          subtitle="Slides a window across features to find LOCAL patterns")

# Convolution visualization (left side)
conv_x = Inches(1.3)
conv_y = Inches(1.3)

add_textbox(slide, conv_x, conv_y, Inches(6), Inches(0.4),
            "How 1D Convolution Works", font_size=18, color=WHITE, bold=True)

# Feature vector representation
feat_y = conv_y + Inches(0.6)
feat_cell_w = Inches(0.65)
feat_cell_h = Inches(0.45)

feature_labels = [
    "OCR_0", "OCR_1", "OCR_2", "...", "UI_29",
    "Vis_0", "Vis_1", "...",
]

for i, label in enumerate(feature_labels):
    fx = conv_x + Inches(0.1) + i * feat_cell_w
    col = DARK_CARD if label != "..." else NUS_BLUE
    if i < 3:
        col = RGBColor(0x1A, 0x5C, 0x3A)  # green tint for filter window on dark bg
    add_rounded_rect(slide, fx, feat_y, feat_cell_w - Inches(0.05),
                     feat_cell_h, col, label, font_size=10, font_color=NUS_ORANGE)
    shape = slide.shapes.add_shape(5, fx, feat_y,
                                   feat_cell_w - Inches(0.05), feat_cell_h)
    shape.fill.background()
    shape.line.color.rgb = LIGHT_GRAY
    shape.line.width = Pt(1)

# Filter bracket
add_rounded_rect(slide, conv_x + Inches(0.1), feat_y + feat_cell_h + Inches(0.05),
                 3 * feat_cell_w - Inches(0.05), Inches(0.35),
                 ACCENT_GREEN, "Filter (size 3)", font_size=11,
                 font_color=WHITE, bold=True)

# Arrow down
add_textbox(slide, conv_x + Inches(1.0), feat_y + feat_cell_h + Inches(0.4),
            Inches(1.0), Inches(0.3), "↓", font_size=16,
            color=ACCENT_GREEN, alignment=PP_ALIGN.CENTER)

# Pattern detected
add_textbox(slide, conv_x + Inches(0.1),
            feat_y + feat_cell_h + Inches(0.7),
            Inches(5), Inches(0.35),
            '[OCR_0, OCR_1, OCR_2] → "pattern in OCR region"',
            font_size=14, color=NUS_ORANGE, bold=True)

# Second filter example at modality boundary
feat_y2 = feat_y + feat_cell_h + Inches(1.3)
boundary_labels = ["...", "UI_29", "Vis_0", "Vis_1", "..."]

for i, label in enumerate(boundary_labels):
    fx = conv_x + Inches(0.1) + i * feat_cell_w
    col = DARK_CARD if label != "..." else NUS_BLUE
    if 1 <= i <= 3:
        col = RGBColor(0x1A, 0x5C, 0x3A)
    add_rounded_rect(slide, fx, feat_y2, feat_cell_w - Inches(0.05),
                     feat_cell_h, col, label, font_size=10, font_color=NUS_ORANGE)
    shape = slide.shapes.add_shape(5, fx, feat_y2,
                                   feat_cell_w - Inches(0.05), feat_cell_h)
    shape.fill.background()
    shape.line.color.rgb = LIGHT_GRAY
    shape.line.width = Pt(1)

add_rounded_rect(slide, conv_x + Inches(0.1) + feat_cell_w,
                 feat_y2 + feat_cell_h + Inches(0.05),
                 3 * feat_cell_w - Inches(0.05), Inches(0.35),
                 ACCENT_GREEN, "Filter (size 3)", font_size=11,
                 font_color=WHITE, bold=True)

add_textbox(slide, conv_x + Inches(0.1),
            feat_y2 + feat_cell_h + Inches(0.45),
            Inches(5.5), Inches(0.35),
            '[UI_29, Vis_0, Vis_1] → "pattern at modality boundary"',
            font_size=14, color=NUS_ORANGE, bold=True)

# Right side
right_x = Inches(7.0)

add_textbox(slide, right_x, Inches(1.3), Inches(5.5), Inches(0.4),
            "Key Insight", font_size=18, color=WHITE, bold=True)

add_rounded_rect(slide, right_x, Inches(1.8), Inches(5.5), Inches(1.2), DARK_CARD)
add_textbox(slide, right_x + Inches(0.2), Inches(1.9), Inches(5.1), Inches(1.0),
            "Finds patterns WITHIN and BETWEEN adjacent features.\n"
            "The filter slides one position at a time, detecting\n"
            "every local pattern in the 150-dim vector.",
            font_size=14, color=NUS_ORANGE)

add_textbox(slide, right_x, Inches(3.3), Inches(5.5), Inches(0.35),
            "Strengths", font_size=16, color=ACCENT_GREEN, bold=True)
add_multiline_textbox(slide, right_x + Inches(0.15), Inches(3.65),
                      Inches(5.2), Inches(0.8),
                      [
                          "•  Detects local feature patterns efficiently",
                          "•  Fewer parameters than MLP",
                          "•  Translation invariant — same pattern detected anywhere",
                      ],
                      font_size=13, color=NUS_ORANGE, spacing=Pt(4))

# What it learns (bottom bar)
add_rounded_rect(slide, Inches(1.3), Inches(5.7), Inches(11.3), Inches(0.8),
                 ACCENT_GREEN)
add_textbox(slide, Inches(1.7), Inches(5.8), Inches(10.5), Inches(0.6),
            'What it learns: "Adjacent features that co-occur — e.g., '
            'high OCR word count + terminal UI = command-line activity"',
            font_size=15, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)


# ============================================================
# SLIDE 7: LSTM
# ============================================================
slide = add_content_slide(prs, "LSTM (Long Short-Term Memory)",
                          subtitle="Processes features as a SEQUENCE, maintaining memory")

# Step-by-step processing (left side)
step_x = Inches(1.3)
step_y = Inches(1.3)
add_textbox(slide, step_x, step_y, Inches(6), Inches(0.4),
            "Sequential Processing", font_size=18, color=WHITE, bold=True)

steps = [
    ("Step 1:", "See OCR[0]=0.82", 'Memory: "probably text editing"', NUS_ORANGE),
    ("Step 50:", "See UI[0]=terminal", 'Memory: "text editing in terminal"', ACCENT_BLUE),
    ("Step 80:", "See Visual=dark", 'Memory: "text in dark terminal"', ACCENT_PURPLE),
    ("Step 150:", "See Interaction=typing", 'Final: "git_operations"', ACCENT_GREEN),
]

sy = step_y + Inches(0.5)
for step_label, feature, memory, col in steps:
    # Step number
    add_rounded_rect(slide, step_x + Inches(0.1), sy,
                     Inches(1.0), Inches(0.5), col,
                     step_label, font_size=13, font_color=WHITE, bold=True)
    # Feature
    add_textbox(slide, step_x + Inches(1.2), sy + Inches(0.05),
                Inches(2.2), Inches(0.4),
                feature, font_size=14, color=NUS_ORANGE, bold=True)
    # Arrow
    add_textbox(slide, step_x + Inches(3.3), sy + Inches(0.05),
                Inches(0.4), Inches(0.4),
                "→", font_size=16, color=LIGHT_GRAY)
    # Memory state
    add_textbox(slide, step_x + Inches(3.6), sy + Inches(0.05),
                Inches(3.0), Inches(0.4),
                memory, font_size=13, color=col, bold=True)

    sy += Inches(0.65)

    # Connecting arrow between steps (except last)
    if step_label != "Step 150:":
        add_textbox(slide, step_x + Inches(0.4), sy - Inches(0.2),
                    Inches(0.4), Inches(0.3),
                    "↓", font_size=14, color=LIGHT_GRAY,
                    alignment=PP_ALIGN.CENTER)

add_textbox(slide, step_x + Inches(0.1), sy + Inches(0.15),
            Inches(6), Inches(0.4),
            "Remembers earlier features when processing later ones",
            font_size=14, color=LIGHT_GRAY, bold=True)

# LSTM cell concept (right side)
cell_x = Inches(7.4)
cell_y = Inches(1.3)

add_textbox(slide, cell_x, cell_y, Inches(5.5), Inches(0.4),
            "LSTM Cell Concept", font_size=18, color=WHITE, bold=True)

# Three gates
gate_w = Inches(5.0)
gate_h = Inches(0.6)
gate_y = cell_y + Inches(0.5)
gates = [
    ("Forget Gate", "What to discard from memory",
     ACCENT_RED),
    ("Input Gate", "What new information to store",
     ACCENT_BLUE),
    ("Output Gate", "What to output from memory",
     ACCENT_GREEN),
]

for gate_name, gate_desc, gate_col in gates:
    add_rounded_rect(slide, cell_x, gate_y, Inches(1.8), gate_h,
                     gate_col, gate_name, font_size=13,
                     font_color=WHITE, bold=True)
    add_textbox(slide, cell_x + Inches(2.0), gate_y + Inches(0.1),
                Inches(3.2), Inches(0.4),
                gate_desc, font_size=13, color=NUS_ORANGE)
    gate_y += gate_h + Inches(0.15)

# Strengths
add_textbox(slide, cell_x, gate_y + Inches(0.3), Inches(5.5), Inches(0.35),
            "Strengths", font_size=16, color=NUS_ORANGE, bold=True)
add_multiline_textbox(slide, cell_x + Inches(0.15), gate_y + Inches(0.65),
                      Inches(5.2), Inches(0.8),
                      [
                          "•  Captures long-range dependencies",
                          "•  Learns cross-modality relationships",
                          "•  Selective memory — keeps what matters",
                      ],
                      font_size=13, color=NUS_ORANGE, spacing=Pt(4))

# What it learns (bottom bar)
add_rounded_rect(slide, Inches(1.3), Inches(5.7), Inches(11.3), Inches(0.8),
                 NUS_ORANGE)
add_textbox(slide, Inches(1.7), Inches(5.8), Inches(10.5), Inches(0.6),
            'What it learns: "How features across ALL modalities relate '
            'to each other sequentially"',
            font_size=15, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)


# ============================================================
# SLIDE 8: Transformer
# ============================================================
slide = add_content_slide(prs, "Transformer (Self-Attention)",
                          subtitle="Uses SELF-ATTENTION to learn which features matter most")

# Attention concept visualization (left side)
att_x = Inches(1.3)
att_y = Inches(1.3)

add_textbox(slide, att_x, att_y, Inches(6), Inches(0.4),
            "Self-Attention Mechanism", font_size=18, color=WHITE, bold=True)

# Feature pairs with attention
pairs = [
    ('Feature 0 (OCR: "git")', 'Feature 120 (Interaction: typing)',
     "attends to", ACCENT_PURPLE),
    ('Feature 50 (UI: terminal)', 'Feature 80 (Visual: dark)',
     "attends to", ACCENT_BLUE),
]

py = att_y + Inches(0.6)
for left_feat, right_feat, relation, col in pairs:
    # Left feature box
    add_rounded_rect(slide, att_x + Inches(0.1), py, Inches(2.8), Inches(0.55),
                     col, left_feat, font_size=11, font_color=WHITE, bold=False)
    # Arrow with label
    add_textbox(slide, att_x + Inches(3.0), py + Inches(0.1),
                Inches(1.6), Inches(0.35),
                "← " + relation + " →",
                font_size=12, color=col, bold=True, alignment=PP_ALIGN.CENTER)
    # Right feature box
    add_rounded_rect(slide, att_x + Inches(4.7), py, Inches(3.0), Inches(0.55),
                     col, right_feat, font_size=11, font_color=WHITE, bold=False)

    py += Inches(0.75)

# Key property
add_rounded_rect(slide, att_x + Inches(0.1), py + Inches(0.1),
                 Inches(7.5), Inches(0.55), DARK_CARD)
add_textbox(slide, att_x + Inches(0.3), py + Inches(0.15),
            Inches(7.1), Inches(0.45),
            "Can connect ANY two features regardless of position — "
            "no sequential processing needed",
            font_size=14, color=NUS_ORANGE, bold=True)

# ChatGPT note
add_textbox(slide, att_x + Inches(0.1), py + Inches(0.8),
            Inches(7.5), Inches(0.4),
            "Same architecture behind ChatGPT, adapted for classification",
            font_size=14, color=LIGHT_GRAY, bold=True)

# Right side: Strengths
right_x = Inches(9.0)

add_textbox(slide, right_x, Inches(1.3), Inches(4.0), Inches(0.4),
            "Strengths", font_size=18, color=ACCENT_PURPLE, bold=True)

add_multiline_textbox(slide, right_x + Inches(0.15), Inches(1.8),
                      Inches(3.8), Inches(1.5),
                      [
                          "•  Learns global relationships",
                          "•  Fully parallelizable (fast GPU)",
                          "•  Attention weights are",
                          "   interpretable",
                          "•  Scales well with data",
                      ],
                      font_size=14, color=NUS_ORANGE, spacing=Pt(4))

# Attention weight visualization hint
add_textbox(slide, right_x, Inches(3.6), Inches(4.0), Inches(0.35),
            "How Attention Works", font_size=16, color=WHITE, bold=True)
add_rounded_rect(slide, right_x, Inches(4.0), Inches(4.0), Inches(1.5), DARK_CARD)
add_multiline_textbox(slide, right_x + Inches(0.15), Inches(4.1),
                      Inches(3.7), Inches(1.3),
                      [
                          "For each feature, compute:",
                          {"text": "  Query: What am I looking for?", "size": 12, "color": ACCENT_PURPLE},
                          {"text": "  Key: What do I contain?", "size": 12, "color": ACCENT_BLUE},
                          {"text": "  Value: What info do I carry?", "size": 12, "color": ACCENT_GREEN},
                          "",
                          "Match queries to keys → attention",
                      ],
                      font_size=12, color=NUS_ORANGE, spacing=Pt(2))

# What it learns (bottom bar)
add_rounded_rect(slide, Inches(1.3), Inches(5.9), Inches(11.3), Inches(0.7),
                 ACCENT_PURPLE)
add_textbox(slide, Inches(1.7), Inches(6.0), Inches(10.5), Inches(0.5),
            'What it learns: "Which feature PAIRS are most informative '
            'for each class"',
            font_size=15, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)


# ============================================================
# SLIDE 9: Tier 2 Comparison
# ============================================================
slide = add_content_slide(prs, "Tier 2: Model Comparison")

# Results table
table_x = Inches(1.3)
table_y = Inches(1.3)
table_w = Inches(11.0)
cols = 5
rows = 5  # header + 4 models

col_widths = [Inches(2.3), Inches(2.2), Inches(2.2), Inches(2.2), Inches(2.1)]

table_shape = slide.shapes.add_table(rows, cols, table_x, table_y,
                                     table_w, Inches(2.8))
table = table_shape.table

# Set column widths
for i, w in enumerate(col_widths):
    table.columns[i].width = w

# Header row
headers = ["Model", "Accuracy", "F1 Score", "Latency", "Training Time"]
for i, h in enumerate(headers):
    cell = table.cell(0, i)
    cell.text = h
    p = cell.text_frame.paragraphs[0]
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Calibri"
    p.alignment = PP_ALIGN.CENTER
    cell.fill.solid()
    cell.fill.fore_color.rgb = ACCENT_PURPLE

# Data rows
data = [
    ["MLP", "72.4%", "0.71", "4ms", "~2 min"],
    ["CNN-1D", "70.8%", "0.69", "5ms", "~3 min"],
    ["LSTM", "69.5%", "0.68", "8ms", "~5 min"],
    ["Transformer", "71.2%", "0.70", "6ms", "~4 min"],
]

model_colors = [ACCENT_BLUE, ACCENT_GREEN, NUS_ORANGE, ACCENT_PURPLE]

for r, row_data in enumerate(data):
    for c, val in enumerate(row_data):
        cell = table.cell(r + 1, c)
        cell.text = val
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(15)
        p.font.name = "Calibri"
        p.alignment = PP_ALIGN.CENTER
        if c == 0:
            p.font.bold = True
            p.font.color.rgb = model_colors[r]
        else:
            p.font.color.rgb = NUS_ORANGE
        # Alternate row background
        if r % 2 == 0:
            cell.fill.solid()
            cell.fill.fore_color.rgb = DARK_CARD
        else:
            cell.fill.solid()
            cell.fill.fore_color.rgb = NUS_BLUE

# Key insight box
insight_y = Inches(4.4)
add_rounded_rect(slide, Inches(1.3), insight_y, Inches(11.0), Inches(1.6),
                 DARK_CARD)

add_textbox(slide, Inches(1.7), insight_y + Inches(0.15),
            Inches(10.2), Inches(0.4),
            "Key Insight", font_size=20, color=ACCENT_PURPLE, bold=True)

add_textbox(slide, Inches(1.7), insight_y + Inches(0.6),
            Inches(10.2), Inches(0.9),
            "Deep learning models haven't outperformed classical ML here because "
            "the 150-dim features are already well-engineered. Deep learning "
            "shines with raw/complex input (images, raw text, audio). "
            "With pre-extracted features, simpler models can be equally effective.",
            font_size=16, color=NUS_ORANGE)

# Decorative note
add_textbox(slide, Inches(1.3), Inches(6.3), Inches(11.0), Inches(0.4),
            "* Results based on CUA-Suite dataset with 9 screen activity classes  |  "
            "All models trained on identical 150-dim feature vectors",
            font_size=12, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)


# ============================================================
# Save
# ============================================================
out_path = "/Users/muneeswaranm/Projects/PRS/PatternRecognSystem-Video2Text/presentation/Video2Knowledge_Tier2_DeepLearning.pptx"
prs.save(out_path)
print(f"Saved: {out_path}")
print(f"Slides: {len(prs.slides)}")
