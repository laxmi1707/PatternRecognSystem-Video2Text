from __future__ import annotations

import io
import logging
from datetime import datetime
from pathlib import Path

import cv2

logger = logging.getLogger(__name__)


def generate_report(
    job_id: int,
    video_path: str | None,
    segments: list[dict],
    task_instruction: str,
    platform: str,
    sop_steps: list[dict],
    anomalies_by_segment: dict[int, list] | None = None,
    keyframe_interval: float = 5.0,
) -> bytes:
    """Build a Word (.docx) report and return it as bytes."""
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError as e:
        raise RuntimeError("python-docx is not installed") from e

    doc = Document()

    # --- Title ---
    title_para = doc.add_heading(task_instruction, level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph(
        f"Video Analysis Report  |  Platform: {platform}  |  Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    )
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    sub.runs[0].font.size = Pt(10)

    doc.add_paragraph()

    # --- Anomaly Summary Banner (if any) ---
    if anomalies_by_segment:
        all_anomalies = [a for lst in anomalies_by_segment.values() for a in lst]
        critical = [a for a in all_anomalies if a.severity == "critical"]
        warnings = [a for a in all_anomalies if a.severity == "warning"]
        banner = doc.add_paragraph()
        banner_run = banner.add_run(
            f"⚠  ANOMALY ALERT: {len(all_anomalies)} anomaly(ies) detected"
            f" — {len(critical)} critical, {len(warnings)} warning(s)"
        )
        banner_run.font.bold = True
        banner_run.font.size = Pt(12)
        banner_run.font.color.rgb = RGBColor(0xDC, 0x26, 0x26)  # red

        # Summary table
        if all_anomalies:
            tbl = doc.add_table(rows=1, cols=3)
            tbl.style = "Light Shading Accent 2"
            hdr = tbl.rows[0].cells
            for i, h in enumerate(["Severity", "Category", "Detected Text"]):
                hdr[i].text = h
                hdr[i].paragraphs[0].runs[0].font.bold = True
            for seg_idx, a_list in sorted(anomalies_by_segment.items()):
                for a in a_list:
                    row = tbl.add_row().cells
                    row[0].text = a.severity.upper()
                    row[1].text = a.category
                    row[2].text = a.matched_text
        doc.add_paragraph()

    # --- Section 1: Keyframes ---
    doc.add_heading("1. Keyframes", level=1)
    doc.add_paragraph(f"One frame extracted every {int(keyframe_interval)} seconds.")

    if video_path and Path(video_path).exists():
        cap = cv2.VideoCapture(video_path)
        try:
            fps = cap.get(cv2.CAP_PROP_FPS) or 30.0
            total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
            duration = total_frames / fps
            t = 0.0
            while t < duration:
                cap.set(cv2.CAP_PROP_POS_MSEC, t * 1000)
                ret, frame = cap.read()
                if not ret:
                    t += keyframe_interval
                    continue
                # Resize to max width 960 to keep doc size reasonable
                h, w = frame.shape[:2]
                if w > 960:
                    scale = 960 / w
                    frame = cv2.resize(frame, (960, int(h * scale)))
                # Encode to PNG bytes
                ok, buf = cv2.imencode(".png", frame)
                if ok:
                    img_stream = io.BytesIO(buf.tobytes())
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    run = p.add_run(f"t = {t:.1f}s")
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
                    doc.add_picture(img_stream, width=Inches(6.0))
                t += keyframe_interval
        finally:
            cap.release()
    else:
        doc.add_paragraph("(Video file not available — keyframes could not be extracted.)")

    doc.add_page_break()

    # --- Section 2: Classification Results ---
    doc.add_heading("2. Classification Results", level=1)
    doc.add_paragraph(f"Total segments analysed: {len(segments)}")

    if segments:
        table = doc.add_table(rows=1, cols=5)
        table.style = "Light Shading Accent 1"
        hdr = table.rows[0].cells
        for i, h in enumerate(["#", "Time Range", "Activity Label", "Confidence", "Model"]):
            hdr[i].text = h
            hdr[i].paragraphs[0].runs[0].font.bold = True

        for i, seg in enumerate(segments, 1):
            row = table.add_row().cells
            start = seg.get("start_time", 0)
            end = seg.get("end_time", 0)
            label = seg.get("predicted_label", "other")
            conf = seg.get("confidence", 0)
            model = seg.get("model_name", "")
            row[0].text = str(i)
            row[1].text = f"{start:.1f}s – {end:.1f}s"
            row[2].text = label.replace("_", " ").title()
            row[3].text = f"{conf * 100:.1f}%"
            row[4].text = model

    doc.add_page_break()

    # --- Section 3: SOP Steps ---
    doc.add_heading("3. Standard Operating Procedure", level=1)

    if sop_steps:
        for step in sop_steps:
            num = step.get("number", "")
            title = step.get("title", "")
            description = step.get("description", "")
            time_range = step.get("time_range", "")

            # Segment index is 0-based; step number is 1-based
            seg_anomalies = (anomalies_by_segment or {}).get(int(num) - 1, [])
            prefix = "⚠ " if seg_anomalies else ""
            doc.add_heading(f"Step {num}: {prefix}{title}", level=2)
            if time_range:
                p = doc.add_paragraph()
                run = p.add_run(f"Time: {time_range}")
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
            if description:
                doc.add_paragraph(description)
            # Inline anomaly warnings for this step
            for a in seg_anomalies:
                warn_p = doc.add_paragraph()
                warn_run = warn_p.add_run(
                    f"⚠ [{a.severity.upper()}] {a.category}: \"{a.matched_text}\" — {a.description}"
                )
                warn_run.font.bold = True
                warn_run.font.color.rgb = RGBColor(0xDC, 0x26, 0x26)
    else:
        doc.add_paragraph("(No SOP steps available.)")

    # Save to bytes
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
